"""DOCX parser using python-docx."""

from __future__ import annotations

from pathlib import Path

import docx

from app.ingestion.parsers._shared import ParsedDocument, _normalize_text


def extract_docx(file_path: str | Path) -> ParsedDocument:
    """Extract text and paragraph metadata from a DOCX file.

    Returns
    -------
    ParsedDocument
        With ``text`` as the full document text and ``metadata["paragraph_count"]``.
    """
    doc = docx.Document(str(file_path))
    paragraphs: list[str] = []
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            paragraphs.append(_normalize_text(t))
    return ParsedDocument(
        text="\n\n".join(paragraphs),
        metadata={"paragraph_count": len(paragraphs)},
    )
